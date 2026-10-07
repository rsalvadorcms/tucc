import streamlit as st
import pandas as pd
import psycopg2
import psycopg2.extras
import io
import urllib.parse
import string
import os
import re
import shutil
import base64
from datetime import datetime, date, timedelta, time

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage

# Optional import for PDF rendering
try:
    from reportlab.lib.pagesizes import A3, A4, portrait
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

# Set page configurations with native default theme formatting
st.set_page_config(page_title="Office Operations Portal", layout="wide")

LOGO1_PATH = "logo.png"
LOGO2_PATH = "logo2.png"
BACKUP_DIR = "backups"
NEWS_DIR = "site_news_uploads"

# Ensure directories exist
if not os.path.exists(BACKUP_DIR):
    os.makedirs(BACKUP_DIR)
if not os.path.exists(NEWS_DIR):
    os.makedirs(NEWS_DIR)

# ==============================================================================
# ⚙ 1. HELPER FUNCTIONS & POSTGRESQL DATABASE ENGINE
# ==============================================================================
# @st.cache_resource
def get_db_connection():
    if "postgres" in st.secrets:
        conn = psycopg2.connect(**st.secrets["postgres"], cursor_factory=psycopg2.extras.DictCursor)
    else:
        conn = psycopg2.connect(
            dbname=os.environ.get("PG_DATABASE", "postgres"),
            user=os.environ.get("PG_USER", "postgres"),
            password=os.environ.get("PG_PASSWORD", "password"),
            host=os.environ.get("PG_HOST", "localhost"),
            port=os.environ.get("PG_PORT", "5432"),
            cursor_factory=psycopg2.extras.DictCursor
        )
    return conn
# Cached Data Query Functions with st.cache_data
@st.cache_data
def get_home_unrolled_data():
    conn = get_db_connection()
    df = pd.read_sql_query('''
        SELECT tg.group_name AS "Car Group", c.vehicle AS "Vehicle Model", c.plate_number AS "Plate Number", 
               c.color AS "Color", tg.driver_name AS "Driver Name", fd.driver_mobile AS "Contact Number", 
               tg.etd_1 AS "ETD 1 (From)", tg.etd_2 AS "ETD 2 (To)", tp.passengers AS "Passenger Name"
        FROM transit_groups tg
        LEFT JOIN cars c ON tg.group_name = c.car_name
        LEFT JOIN fleet_drivers fd ON tg.driver_name = fd.driver_name
        LEFT JOIN transit_passengers tp ON tg.group_name = tp.group_name
        ORDER BY tg.group_name
    ''', conn)
    return df

@st.cache_data
def get_shuttle_raw_data():
    conn = get_db_connection()
    df = pd.read_sql_query('''
        SELECT dt.id, dt.transit_date_start, dt.transit_date_end, 
               dt.trip AS trip_code,
               dt.requested_by, dt.group_name, c.plate_number, c.vehicle, 
               dt.location_from, dt.location_to, dt.etd_1, dt.etd_2, 
               tg.driver_name
        FROM daily_transit dt 
        LEFT JOIN cars c ON dt.group_name = c.car_name
        LEFT JOIN transit_groups tg ON dt.group_name = tg.group_name
        WHERE dt.trip IN ('Trip A', 'Trip B', 'Trip C')
        ORDER BY dt.etd_1 ASC, dt.id DESC
    ''', conn)
    return df

@st.cache_data
def get_holidays_list():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT holiday_date FROM holidays", conn)
    return df['holiday_date'].tolist()

@st.cache_data
def get_employees_list():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT emp_name FROM users WHERE emp_name IS NOT NULL AND emp_name != ''", conn)
    return df['emp_name'].tolist()

@st.cache_data
def get_cars_list():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT car_name FROM cars", conn)
    return df['car_name'].tolist()

@st.cache_data
def get_drivers_list():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT driver_name FROM fleet_drivers", conn)
    return df['driver_name'].tolist()

@st.cache_data
def get_trips_list():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT trip FROM trips", conn)
    return df['trip'].tolist()

@st.cache_data
def get_transit_groups_list():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT group_name FROM transit_groups", conn)
    return df['group_name'].tolist()

@st.cache_data
def get_meeting_rooms_df():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT room_number AS Room_Number, room_name AS Room_Name, capacity AS Capacity, location AS Location FROM meeting_rooms", conn)
    return df

@st.cache_data
def get_room_numbers_list():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT room_number FROM meeting_rooms", conn)
    return df['room_number'].tolist()

@st.cache_data
def get_overtime_requests_df(role, username):
    conn = get_db_connection()
    if role in ["Admin", "Owner"]:
        df = pd.read_sql_query("SELECT id, username, emp_name AS Employee_Name, ot_date AS Date, start_time AS Start_Time, end_time AS End_Time, needs_transport AS Needs_Transport, origin AS Origin, destination AS Destination, departure_time AS Departure_Time, return_time AS Return_Time FROM overtime_requests", conn)
    else:
        df = pd.read_sql_query("SELECT id, username, emp_name AS Employee_Name, ot_date AS Date, start_time AS Start_Time, end_time AS End_Time, needs_transport AS Needs_Transport, origin AS Origin, destination AS Destination, departure_time AS Departure_Time, return_time AS Return_Time FROM overtime_requests WHERE username = %s", conn, params=[username])
    return df

@st.cache_data
def get_daily_transit_df():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT id, transit_date_start AS Start, transit_date_end AS End, group_name AS Car_Group, requested_by AS Requested_By, etd_1 AS ETD_1, etd_2 AS ETD_2, location_from AS From, location_to AS To, daily AS Daily, trip AS Trip FROM daily_transit ORDER BY id DESC", conn)
    return df

@st.cache_data
def get_room_bookings_df():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT id, room_number AS Room, booked_by AS Booked_By, booking_date AS Date, start_time AS Start, end_time AS End, is_recurring AS Recurring, recurrence_end_date AS Recurrence_End FROM room_bookings ORDER BY id DESC", conn)
    return df

@st.cache_data
def get_site_news_df():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT id, title, filename, file_path, uploaded_by, upload_date FROM site_news ORDER BY id DESC", conn)
    return df

@st.dialog("Data Transaction Status")
def show_transaction_dialog(title_text: str, message_text: str, status_type: str = "success"):
    """Displays transaction outcome in a pop-up modal message box."""
    if status_type == "success":
        st.success(f"### {title_text}")
    elif status_type == "error":
        st.error(f"### {title_text}")
    else:
        st.info(f"### {title_text}")
        
    st.write(message_text)
    
    if st.button("OK", type="primary", use_container_width=True):
        if 'tx_dialog_title' in st.session_state:
            del st.session_state['tx_dialog_title']
        if 'tx_dialog_msg' in st.session_state:
            del st.session_state['tx_dialog_msg']
        if 'tx_dialog_type' in st.session_state:
            del st.session_state['tx_dialog_type']
        st.rerun()

def trigger_transaction_dialog():
    """Renders the modal popup box if status data exists in session state."""
    if 'tx_dialog_title' in st.session_state and st.session_state.tx_dialog_title:
        show_transaction_dialog(
            st.session_state.tx_dialog_title,
            st.session_state.get('tx_dialog_msg', ''),
            st.session_state.get('tx_dialog_type', 'success')
        )

def set_transaction_dialog(title: str, message: str, status_type: str = "success"):
    """Stores dialog notification parameters into session state."""
    st.session_state.tx_dialog_title = title
    st.session_state.tx_dialog_msg = message
    st.session_state.tx_dialog_type = status_type

def perform_manual_backup():
    """Creates a timestamped manual backup export of ALL database tables to Excel."""
    if not os.path.exists(BACKUP_DIR):
        os.makedirs(BACKUP_DIR)
        
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
    try:
        conn = get_db_connection()
        backup_excel_path = os.path.join(BACKUP_DIR, f"backup_data_{timestamp}.xlsx")
        with pd.ExcelWriter(backup_excel_path, engine='openpyxl') as writer:
            pd.read_sql_query("SELECT * FROM daily_transit", conn).to_excel(writer, index=False, sheet_name="daily_transit")
            pd.read_sql_query("SELECT * FROM overtime_requests", conn).to_excel(writer, index=False, sheet_name="overtime_requests")
            pd.read_sql_query("SELECT * FROM room_bookings", conn).to_excel(writer, index=False, sheet_name="room_bookings")
            pd.read_sql_query("SELECT * FROM users", conn).to_excel(writer, index=False, sheet_name="users")
            pd.read_sql_query("SELECT * FROM cars", conn).to_excel(writer, index=False, sheet_name="cars")
            pd.read_sql_query("SELECT * FROM trips", conn).to_excel(writer, index=False, sheet_name="trips")
            pd.read_sql_query("SELECT * FROM site_news", conn).to_excel(writer, index=False, sheet_name="site_news")
            pd.read_sql_query("SELECT * FROM transit_passengers", conn).to_excel(writer, index=False, sheet_name="transit_passengers")
            pd.read_sql_query("SELECT * FROM holidays", conn).to_excel(writer, index=False, sheet_name="holidays")
            pd.read_sql_query("SELECT * FROM meeting_rooms", conn).to_excel(writer, index=False, sheet_name="meeting_rooms")
            pd.read_sql_query("SELECT * FROM fleet_drivers", conn).to_excel(writer, index=False, sheet_name="fleet_drivers")
            pd.read_sql_query("SELECT * FROM transit_groups", conn).to_excel(writer, index=False, sheet_name="transit_groups")
        return True, "PostgreSQL Backup Export", backup_excel_path
    except Exception as e:
        return False, str(e), ""

def format_military_time(input_str: str) -> str:
    if not input_str:
        return ""
    clean_digits = re.sub(r'\D', '', str(input_str).strip())
    if len(clean_digits) == 3:
        clean_digits = "0" + clean_digits
    elif len(clean_digits) == 1:
        clean_digits = "0" + clean_digits + "00"
    elif len(clean_digits) == 2:
        clean_digits = clean_digits + "00"
    if len(clean_digits) == 4:
        hours = int(clean_digits[:2])
        minutes = int(clean_digits[2:])
        if 0 <= hours <= 23 and 0 <= minutes <= 59:
            return f"{hours:02d}:{minutes:02d}"
    return None

def export_df_to_excel(df, sheet_name="Data"):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
    return buffer.getvalue()

def export_overtime_summary_excel(df, title_text="Overtime Schedule"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Overtime Summary"
    ws.views.sheetView[0].showGridLines = True
    
    font_title = Font(name="Calibri", size=12, bold=True, color="1F4E78")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF')
    )
    
    ws.merge_cells("A1:E1")
    t_cell = ws["A1"]
    t_cell.value = title_text
    t_cell.font = font_title
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25
    
    headers = list(df.columns)
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[3].height = 20
    
    for r_idx, row in enumerate(df.iterrows(), start=4):
        for c_idx, val in enumerate(row[1], start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value="" if pd.isna(val) else val)
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    for col_idx in range(1, ws.max_column + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        max_len = max((len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(1, ws.max_row + 1)), default=0)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 18)
        
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def export_overtime_summary_pdf(df, title_text="Overtime Schedule"):
    if not HAS_REPORTLAB:
        return None
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=portrait(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=13, leading=16, alignment=1, textColor=colors.HexColor('#1F4E78'))
    cell_style = ParagraphStyle('CellSt', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=11, alignment=1)
    header_style = ParagraphStyle('HeadSt', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9, leading=11, alignment=1, textColor=colors.white)
    
    story = [Paragraph(title_text, title_style), Spacer(1, 15)]
    
    headers = list(df.columns)
    table_data = [[Paragraph(h, header_style) for h in headers]]
    
    for _, row in df.iterrows():
        row_cells = [Paragraph(str(val if pd.notna(val) else ""), cell_style) for val in row]
        table_data.append(row_cells)
        
    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E78')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#BFBFBF')),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def export_custom_batam_excel(detailed_df, effective_date_str=""):
    wb = openpyxl.Workbook()
    font_title = Font(name="Calibri", size=13, bold=True, color="1F4E78")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold_label = Font(name="Calibri", size=10, bold=True, color="000000")
    font_data = Font(name="Calibri", size=10)
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF')
    )

    ws2 = wb.active
    ws2.title = "Detailed Allocations"
    ws2.views.sheetView[0].showGridLines = True
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A3
    ws2.page_setup.orientation = ws2.ORIENTATION_PORTRAIT
    ws2.print_title_rows = '1:6'
    
    for r in range(1, 6):
        for c in range(1, 7):
            ws2.cell(row=r, column=c).fill = white_fill

    ws2.row_dimensions[2].height = 20
    ws2.row_dimensions[3].height = 20
    ws2.row_dimensions[4].height = 20

    if effective_date_str:
        cell_e5 = ws2.cell(row=5, column=5, value="Effective Date:")
        cell_e5.font = font_bold_label
        cell_e5.alignment = Alignment(horizontal="right", vertical="center")
        
        cell_f5 = ws2.cell(row=5, column=6, value=effective_date_str)
        cell_f5.font = font_bold_label
        cell_f5.alignment = Alignment(horizontal="center", vertical="center")

    start_row = 6
    ws2.merge_cells("B2:D4")
    title_cell = ws2["B2"]
    title_cell.value = "Daily Transportation Arrangement - Passenger list\nTUCC PROJECT - BATAM MODULE YARD [MD-1 & MD-4]\nJOB CODE : 0 - 0847 - 00 - 0001"
    title_cell.font = font_title
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    if os.path.exists(LOGO1_PATH):
        try:
            img1_det = OpenpyxlImage(LOGO1_PATH)
            img1_det.width = 110; img1_det.height = 50
            ws2.add_image(img1_det, "A2")
        except Exception:
            pass

    if os.path.exists(LOGO2_PATH):
        try:
            img2_det = OpenpyxlImage(LOGO2_PATH)
            img2_det.width = 130; img2_det.height = 50
            ws2.add_image(img2_det, "F2")
        except Exception:
            pass

    det_headers = ["Vehicle Description", "Driver Name", "Contact Number", "Passenger", "ETD 1", "ETD 2"]
    for col_idx, h_title in enumerate(det_headers, start=1):
        cell = ws2.cell(row=start_row, column=col_idx, value=h_title)
        cell.font = font_header; cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    data_start = start_row + 1
    for r_idx, row in enumerate(detailed_df.iterrows(), start=data_start):
        row_data = row[1]
        v_model = str(row_data.get("Vehicle Model", "") or "Standard Vehicle").strip()
        p_num = str(row_data.get("Plate Number", "") or "N/A").strip()
        v_color = str(row_data.get("Color", "") or "Black").strip()
        v_desc_3lines = f"{v_model}\n{p_num}\nColor : {v_color}"
            
        d_name = row_data.get("Driver Name", "")
        c_num = row_data.get("Contact Number", "")
        p_name = row_data.get("Passenger Name", "")
        etd1 = row_data.get("ETD 1 (From)", "")
        etd2 = row_data.get("ETD 2 (To)", "")
        
        ordered_vals = [v_desc_3lines, d_name, c_num, p_name, etd1, etd2]
        for c_idx, val in enumerate(ordered_vals, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value="" if pd.isna(val) else val)
            cell.font = font_data; cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if not detailed_df.empty:
        current_grp = None
        grp_start = data_start
        tot_rows = len(detailed_df)
        for idx, grp_val in enumerate(detailed_df['Car Group'].values, start=data_start):
            if grp_val != current_grp:
                if current_grp is not None and (idx - 1) > grp_start:
                    for merge_col in [1, 2, 3, 5, 6]:
                        ws2.merge_cells(start_row=grp_start, start_column=merge_col, end_row=idx - 1, end_column=merge_col)
                        ws2.cell(row=grp_start, column=merge_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                current_grp = grp_val
                grp_start = idx
        if current_grp is not None and (data_start + tot_rows - 1) > grp_start:
            for merge_col in [1, 2, 3, 5, 6]:
                ws2.merge_cells(start_row=grp_start, start_column=merge_col, end_row=data_start + tot_rows - 1, end_column=merge_col)
                ws2.cell(row=grp_start, column=merge_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col in ws2.columns:
        col_letter = col[0].column_letter
        if col_letter == 'B':
            ws2.column_dimensions['B'].width = 12
        elif col_letter == 'C':
            ws2.column_dimensions['C'].width = 18
        else:
            max_len = max((len(str(cell.value or '')) for cell in col), default=0)
            ws2.column_dimensions[col_letter].width = max(max_len + 4, 18)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def export_custom_batam_pdf(detailed_df, effective_date_str=""):
    if not HAS_REPORTLAB:
        return None
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=portrait(A3), rightMargin=25, leftMargin=25, topMargin=25, bottomMargin=25)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=15, leading=19, alignment=1, textColor=colors.HexColor('#1F4E78'))
    subtitle_style = ParagraphStyle('DocSubtitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=14, alignment=1, textColor=colors.HexColor('#333333'))
    eff_date_style = ParagraphStyle('EffDateStyle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=12, alignment=2, textColor=colors.HexColor('#000000'))
    cell_style = ParagraphStyle('CellText', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=11, alignment=1)
    header_style = ParagraphStyle('HeaderStyle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=12, alignment=1, textColor=colors.white)
    
    story = []
    title_p = Paragraph("Daily Transportation Arrangement - Passenger list", title_style)
    sub1_p = Paragraph("TUCC PROJECT - BATAM MODULE YARD [MD-1 & MD-4]", subtitle_style)
    sub2_p = Paragraph("JOB CODE : 0 - 0847 - 00 - 0001", subtitle_style)
    header_box = [title_p, Spacer(1, 3), sub1_p, Spacer(1, 2), sub2_p]
    
    img1_elem = RLImage(LOGO1_PATH, width=100, height=45) if os.path.exists(LOGO1_PATH) else ""
    img2_elem = RLImage(LOGO2_PATH, width=120, height=45) if os.path.exists(LOGO2_PATH) else ""
    
    top_table = Table([[img1_elem, header_box, img2_elem]], colWidths=[110, 560, 110])
    top_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('ALIGN', (0,0), (0,0), 'LEFT'), ('ALIGN', (1,0), (1,0), 'CENTER'), ('ALIGN', (2,0), (2,0), 'RIGHT')]))
    story.append(top_table)
    story.append(Spacer(1, 10))
    
    if effective_date_str:
        eff_p = Paragraph(f"<b>Effective Date:</b> {effective_date_str}", eff_date_style)
        eff_table = Table([[Paragraph("", cell_style), eff_p]], colWidths=[550, 230])
        eff_table.setStyle(TableStyle([('ALIGN', (1,0), (1,0), 'RIGHT'), ('VALIGN', (0,0), (-1,-1), 'MIDDLE')]))
        story.append(eff_table)
        story.append(Spacer(1, 8))
    
    headers = ["Vehicle Description", "Driver Name", "Contact Number", "Passenger", "ETD 1", "ETD 2"]
    table_data = [[Paragraph(h, header_style) for h in headers]]
    table_spans = []
    current_grp = None
    grp_start = 1
    
    for idx, row_data in enumerate(detailed_df.iterrows(), start=1):
        r_val = row_data[1]
        grp_name = r_val.get("Car Group", "")
        v_model = str(r_val.get("Vehicle Model", "") or "Standard Vehicle").strip()
        p_num = str(r_val.get("Plate Number", "") or "N/A").strip()
        v_color = str(r_val.get("Color", "") or "Black").strip()
        v_desc_html = f"{v_model}<br/>{p_num}<br/>Color : {v_color}"
        
        row_cells = [
            Paragraph(v_desc_html, cell_style),
            Paragraph(str(r_val.get("Driver Name", "") or ""), cell_style),
            Paragraph(str(r_val.get("Contact Number", "") or ""), cell_style),
            Paragraph(str(r_val.get("Passenger Name", "") or ""), cell_style),
            Paragraph(str(r_val.get("ETD 1 (From)", "") or ""), cell_style),
            Paragraph(str(r_val.get("ETD 2 (To)", "") or ""), cell_style)
        ]
        table_data.append(row_cells)
        if grp_name != current_grp:
            if current_grp is not None and (idx - 1) > grp_start:
                for col_i in [0, 1, 2, 4, 5]:
                    table_spans.append(('SPAN', (col_i, grp_start), (col_i, idx - 1)))
            current_grp = grp_name
            grp_start = idx

    tot_rows = len(detailed_df)
    if current_grp is not None and tot_rows > grp_start:
        for col_i in [0, 1, 2, 4, 5]:
            table_spans.append(('SPAN', (col_i, grp_start), (col_i, tot_rows)))
            
    main_table = Table(table_data, colWidths=[150, 120, 120, 210, 90, 90], repeatRows=1)
    main_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E78')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'), ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#BFBFBF')),
        ('TOPPADDING', (0,0), (-1,-1), 6), ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ] + table_spans))
    story.append(main_table)
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def export_shuttle_timetable_excel(df, effective_date_str=""):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shuttle Timetable"
    ws.views.sheetView[0].showGridLines = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    
    font_title = Font(name="Calibri", size=13, bold=True, color="1F4E78")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF')
    )

    for r in range(1, 20):
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = white_fill

    ws.merge_cells("B2:F3")
    title_cell = ws["B2"]
    title_cell.value = "JGC SHUTTLE TIMETABLE\nTUCC PROJECT - BATAM MODULE YARD [MD-1 & MD-4]"
    title_cell.font = font_title
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if os.path.exists(LOGO1_PATH):
        try:
            img1 = OpenpyxlImage(LOGO1_PATH)
            img1.width = 110; img1.height = 45
            ws.add_image(img1, "A2")
        except Exception:
            pass

    if os.path.exists(LOGO2_PATH):
        try:
            img2 = OpenpyxlImage(LOGO2_PATH)
            img2.width = 120; img2.height = 45
            ws.add_image(img2, "G2")
        except Exception:
            pass

    ws.merge_cells("A5:A6")
    ws["A5"] = "Days (s)"
    
    ws.merge_cells("B5:B6")
    ws["B5"] = "UNIT"
    
    ws.merge_cells("C5:C6")
    ws["C5"] = "DRIVER"
    
    ws.merge_cells("D5:D6")
    ws["D5"] = "TRIP NO."

    ws.merge_cells("E5:F5")
    ws["E5"] = "ROUTE"
    ws["E6"] = "YARD - 1"
    ws["F6"] = "YARD - 3"

    ws.merge_cells("G5:G6")
    ws["G5"] = "REMARKS"

    for r in [5, 6]:
        ws.row_dimensions[r].height = 22
        for col in range(1, 8):
            cell = ws.cell(row=r, column=col)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

    start_row = 7
    sorted_df = df.sort_values(by="ETD Start") if "ETD Start" in df.columns else df
    
    for idx, row in enumerate(sorted_df.iterrows(), start=start_row):
        r_val = row[1]
        unit_val = str(r_val.get("Unit", "TOYOTA HI-ACE"))
        driver_val = str(r_val.get("Driver", "TBA"))
        trip_no = f"{idx - start_row + 1}st" if (idx - start_row + 1) == 1 else f"{idx - start_row + 1}nd" if (idx - start_row + 1) == 2 else f"{idx - start_row + 1}rd" if (idx - start_row + 1) == 3 else f"{idx - start_row + 1}th"
        etd1 = str(r_val.get("ETD Start", ""))
        etd2 = str(r_val.get("ETD Return", ""))
        remarks = str(r_val.get("Remarks", "DROP-OFF"))

        row_values = ["", unit_val, driver_val, trip_no, etd1, etd2, remarks]
        ws.row_dimensions[idx].height = 20
        for c_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    total_rows = len(sorted_df)
    if total_rows > 0:
        end_row = start_row + total_rows - 1
        ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
        days_cell = ws.cell(row=start_row, column=1)
        days_cell.value = "MONDAY\nTUESDAY\nWEDNESDAY\nTHURSDAY\nFRIDAY\nSATURDAY"
        days_cell.font = font_data
        days_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        days_cell.border = thin_border

    max_unit_len = max([len(str(r.get("Unit", ""))) for _, r in sorted_df.iterrows()] + [4], default=18)
    col_width_b = max(max_unit_len + 4, 22)

    col_widths = {'A': 18, 'B': col_width_b, 'C': 16, 'D': 12, 'E': 14, 'F': 14, 'G': 18}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS users (
            username TEXT PRIMARY KEY, password TEXT NOT NULL,
            role TEXT NOT NULL, email_recipients TEXT, emp_name TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS holidays (
            holiday_date TEXT PRIMARY KEY, description TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS trips (
            trip TEXT PRIMARY KEY, trip_name TEXT NOT NULL
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS overtime_requests (
            id SERIAL PRIMARY KEY, username TEXT, emp_name TEXT,
            ot_date TEXT, start_time TEXT, end_time TEXT, needs_transport TEXT DEFAULT 'Yes',
            origin TEXT, destination TEXT, departure_time TEXT, return_time TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS meeting_rooms (
            room_number TEXT PRIMARY KEY, room_name TEXT, capacity INTEGER, location TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS room_bookings (
            id SERIAL PRIMARY KEY, room_number TEXT, booked_by TEXT,
            booking_date TEXT, start_time TEXT, end_time TEXT, is_recurring TEXT, recurrence_end_date TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS cars (
            car_name TEXT PRIMARY KEY, plate_number TEXT UNIQUE NOT NULL, vehicle TEXT, color TEXT DEFAULT 'Black'
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS fleet_drivers (
            driver_name TEXT PRIMARY KEY, driver_mobile TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS transit_groups (
            id SERIAL PRIMARY KEY, group_name TEXT UNIQUE NOT NULL,
            driver_name TEXT, etd_1 TEXT, etd_2 TEXT,
            FOREIGN KEY (group_name) REFERENCES cars(car_name) ON DELETE CASCADE,
            FOREIGN KEY (driver_name) REFERENCES fleet_drivers(driver_name) ON DELETE SET NULL
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS transit_passengers (
            id SERIAL PRIMARY KEY, group_name TEXT NOT NULL, passengers TEXT NOT NULL,
            FOREIGN KEY (group_name) REFERENCES transit_groups(group_name) ON DELETE CASCADE
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS daily_transit (
            id SERIAL PRIMARY KEY, transit_date_start TEXT NOT NULL, transit_date_end TEXT NOT NULL,
            group_name TEXT DEFAULT 'TBA', requested_by TEXT, etd_1 TEXT, etd_2 TEXT,
            location_from TEXT, location_to TEXT, daily TEXT DEFAULT 'No', trip TEXT DEFAULT 'Trip A'
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS site_news (
            id SERIAL PRIMARY KEY, title TEXT NOT NULL, filename TEXT NOT NULL,
            file_path TEXT NOT NULL, file_type TEXT NOT NULL, uploaded_by TEXT, upload_date TEXT
        )
    ''')
    
    # Verified and created database indexes for faster query performance
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_overtime_date ON overtime_requests(ot_date);')
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_overtime_username ON overtime_requests(username);')
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_daily_transit_dates ON daily_transit(transit_date_start, transit_date_end);')
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_daily_transit_trip ON daily_transit(trip);')
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_room_bookings_date ON room_bookings(booking_date);')
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_transit_passengers_group ON transit_passengers(group_name);')

    cursor.execute("SELECT COUNT(*) FROM users")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO users VALUES (%s, %s, %s, %s, %s)", ('owner', 'owner123', 'Owner', 'owner@company.com', 'System Owner'))
        cursor.execute("INSERT INTO users VALUES (%s, %s, %s, %s, %s)", ('admin', 'admin123', 'Admin', 'admin@company.com', 'System Administrator'))

    cursor.execute("SELECT COUNT(*) FROM trips")
    if cursor.fetchone()[0] == 0:
        for t_code, t_desc in [("Trip A", "Yard to Yard"), ("Trip B", "Sunday Panbil - Wasco - Panbil"), ("Trip C", "Panbil - Destination - Panbil"), ("Trip D", "Overtime Dispatch Route")]:
            cursor.execute("INSERT INTO trips (trip, trip_name) VALUES (%s, %s)", (t_code, t_desc))
        
    cursor.execute("SELECT COUNT(*) FROM meeting_rooms")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO meeting_rooms VALUES (%s, %s, %s, %s)", ('101', 'Boardroom', 15, '1st Floor'))
        cursor.execute("INSERT INTO meeting_rooms VALUES (%s, %s, %s, %s)", ('102', 'Huddle Room Alpha', 6, '2nd Floor'))

    cursor.execute("SELECT COUNT(*) FROM fleet_drivers")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO fleet_drivers VALUES (%s, %s)", ('John Doe', '+628111222333'))
        cursor.execute("INSERT INTO fleet_drivers VALUES (%s, %s)", ('Jane Smith', '+628999888777'))

    cursor.execute("SELECT COUNT(*) FROM cars WHERE car_name = 'TBA'")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO cars (car_name, plate_number, vehicle, color) VALUES (%s, %s, %s, %s)", ('TBA', 'TBA', 'TBA', 'TBA'))

    cursor.execute("SELECT COUNT(*) FROM cars")
    if cursor.fetchone()[0] == 1:
        cursor.execute("INSERT INTO cars VALUES (%s, %s, %s, %s) ON CONFLICT (car_name) DO NOTHING", ('Car A', 'B 1234 ABC', 'Toyota Avanza', 'Black'))
        cursor.execute("INSERT INTO cars VALUES (%s, %s, %s, %s) ON CONFLICT (car_name) DO NOTHING", ('Car B', 'B 5678 XYZ', 'Toyota Innova', 'White'))
        
    conn.commit()
    conn.close()

init_db()

def run_migrations():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT column_name 
        FROM information_schema.columns 
        WHERE table_name = 'cars'
    """)
    columns = [col[0] for col in cursor.fetchall()]
    if "color" not in columns:
        cursor.execute("ALTER TABLE cars ADD COLUMN color TEXT DEFAULT 'Black'")
    conn.commit()
    conn.close()

run_migrations()
trigger_transaction_dialog()

# ==============================================================================
# 🔐 2. AUTHENTICATION USER INTERFACE
# ==============================================================================
if 'logged_in' not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.emp_name = ""

if 'emp_name' not in st.session_state:
    st.session_state.emp_name = st.session_state.get('username', '')

if not st.session_state.logged_in:
    st.title("🏢 Office Operations Portal")
    st.subheader("Login to access scheduling & overtime systems")
    
    with st.form("login_form"):
        username = st.text_input("Username")
        password = st.text_input("Password", type="password")
        submitted = st.form_submit_button("Log In")
        
        if submitted:
            conn = get_db_connection()
            with conn.cursor() as cur:
                cur.execute("SELECT * FROM users WHERE username=%s AND password=%s", (username, password))
                user = cur.fetchone()
            if user:
                st.session_state.logged_in = True
                st.session_state.username = user['username']
                st.session_state.role = user['role']
                st.session_state.emp_name = user['emp_name'] or user['username']
                st.rerun()
            else:
                st.error("Invalid username or password configuration.")
    st.stop()

# ==============================================================================
# 🗂️ 3. MAIN APP CONTROL PANELS & SIDEBAR NAVIGATION (Mobile Optimized & Left-Aligned)
# ==============================================================================
st.sidebar.title(f"👋 Welcome, {st.session_state.emp_name}")
st.sidebar.info(f"Access Level: **{st.session_state.role}**")

st.sidebar.markdown("---")
st.sidebar.subheader("📌 Navigation")

st.sidebar.markdown(
    """
    <style>
    [data-testid="stSidebar"] button {
        display: flex !important;
        justify-content: flex-start !important;
        text-align: left !important;
        width: 100% !important;
        padding-top: 10px !important;
        padding-bottom: 10px !important;
    }
    [data-testid="stSidebar"] button p, 
    [data-testid="stSidebar"] button div,
    [data-testid="stSidebar"] button span {
        text-align: left !important;
        justify-content: flex-start !important;
        width: 100% !important;
    }

    @media (max-width: 768px) {
        .freeze-pane-container {
            max-height: 450px !important;
        }
        h1 {
            font-size: 1.4rem !important;
        }
        h2 {
            font-size: 1.2rem !important;
        }
        h3 {
            font-size: 1.05rem !important;
        }
        .stButton button {
            width: 100% !important;
            margin-bottom: 4px !important;
            font-size: 14px !important;
            padding: 8px 12px !important;
        }
    }
    </style>
    """,
    unsafe_allow_html=True
)

nav_options = [
    "🏠 Daily Transportation Arrangement",
    "📅 Shuttle Timetable",
    "⏰ Overtime & Transport",
    "👥 Transit Groups & Passengers",
    "📅 Daily Transit Dispatch Setup",
    "🏢 Meeting Rooms",
    "📢 Site News",
    "🛠️ System Administration"
]

if 'nav_selection' not in st.session_state:
    st.session_state.nav_selection = nav_options[0]

for opt in nav_options:
    is_active = (st.session_state.nav_selection == opt)
    
    if is_active:
        st.sidebar.markdown(
            f"""
            <div style="background-color: #FFF2CC; border-left: 5px solid #1F4E78; padding: 10px 12px; margin-bottom: 6px; border-radius: 4px; text-align: left;">
                {opt}
            </div>
            """,
            unsafe_allow_html=True
        )
    else:
        if st.sidebar.button(opt, key=f"nav_btn_{opt}", use_container_width=True):
            st.session_state.nav_selection = opt
            st.rerun()

nav_selection = st.session_state.nav_selection

st.sidebar.markdown("---")
if st.sidebar.button("Logout Profile"):
    st.session_state.logged_in = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.emp_name = ""
    st.rerun()

# --- VIEW RENDERERS BASED ON SIDEBAR NAVIGATION ---

# 1. DAILY TRANSPORTATION ARRANGEMENT
if nav_selection == "🏠 Daily Transportation Arrangement":
    st.markdown(
        """
        <style>
        .freeze-pane-container {
            max-height: 700px;
            overflow-y: auto;
            border: 1px solid #BFBFBF;
            border-radius: 5px;
            background-color: white;
            padding: 10px;
        }
        .freeze-pane-container table {
            width: 100%;
            border-collapse: collapse;
            background-color: white;
            color: black;
            font-family: Calibri, sans-serif;
            font-size: 14px;
        }
        .freeze-pane-container th {
            position: sticky;
            top: 0;
            background-color: #1F4E78;
            color: white;
            text-align: center;
            border: 1px solid #BFBFBF;
            padding: 10px;
            z-index: 5;
        }
        .freeze-pane-container td {
            border: 1px solid #BFBFBF;
            text-align: center;
            vertical-align: middle;
            padding: 8px;
            background-color: white;
        }
        </style>
        """,
        unsafe_allow_html=True
    )

    h_col1, h_col2, h_col3 = st.columns([1, 4, 1])
    with h_col1:
        if os.path.exists(LOGO1_PATH):
            st.image(LOGO1_PATH, width=130)
    with h_col2:
        st.markdown(
            """
            <div style="text-align: center;">
                <h3 style="color: #1F4E78; margin-bottom: 0px;">Daily Transportation Arrangement - Passenger list</h3>
                <h4 style="color: #333333; margin-top: 2px; margin-bottom: 2px;">TUCC PROJECT - BATAM MODULE YARD [MD-1 & MD-4]</h4>
                <p style="color: #555555; font-size: 14px; font-weight: bold; margin-top: 0px;">JOB CODE : 0 - 0847 - 00 - 0001</p>
            </div>
            """,
            unsafe_allow_html=True
        )
    with h_col3:
        if os.path.exists(LOGO2_PATH):
            st.image(LOGO2_PATH, width=140)

    st.markdown("---")

    home_unrolled_df = get_home_unrolled_data()

    if not home_unrolled_df.empty:
        display_home_df = home_unrolled_df.copy()
        display_home_df["Vehicle Description"] = display_home_df.apply(
            lambda r: f"{r['Vehicle Model'] or 'Standard Vehicle'}<br>{r['Plate Number'] or 'N/A'}<br>Color : {r['Color'] or 'Black'}",
            axis=1
        )
        
        html_table_rows = ""
        current_group = None
        group_rowspan_counts = display_home_df['Car Group'].value_counts()
        
        rendered_groups = set()
        for _, row in display_home_df.iterrows():
            g_name = row['Car Group']
            v_desc = row['Vehicle Description']
            d_name = row['Driver Name'] or ""
            c_num = row['Contact Number'] or ""
            p_name = row['Passenger Name'] or ""
            etd1 = row['ETD 1 (From)'] or ""
            etd2 = row['ETD 2 (To)'] or ""
            
            span_count = group_rowspan_counts.get(g_name, 1)
            
            html_table_rows += "<tr>"
            if g_name != current_group:
                current_group = g_name
                rendered_groups.clear()
                html_table_rows += f"<td rowspan='{span_count}'><b>{g_name}</b><br><span style='font-size:11px;'>{v_desc}</span></td>"
                html_table_rows += f"<td rowspan='{span_count}'>{d_name}</td>"
                html_table_rows += f"<td rowspan='{span_count}'>{c_num}</td>"
            
            html_table_rows += f"<td>{p_name}</td>"
            
            if g_name not in rendered_groups:
                rendered_groups.add(g_name)
                html_table_rows += f"<td rowspan='{span_count}'>{etd1}</td>"
                html_table_rows += f"<td rowspan='{span_count}'>{etd2}</td>"
            
            html_table_rows += "</tr>"

        freeze_pane_html = f"""
        <div class="freeze-pane-container">
            <table>
                <thead>
                    <tr>
                        <th>Vehicle Description</th>
                        <th>Driver Name</th>
                        <th>Contact Number</th>
                        <th>Passenger</th>
                        <th>ETD 1</th>
                        <th>ETD 2</th>
                    </tr>
                </thead>
                <tbody>
                    {html_table_rows}
                </tbody>
            </table>
        </div>
        """
        st.markdown(freeze_pane_html, unsafe_allow_html=True)
    else:
        st.info("No transit groups or passenger assignments configured yet. Go to 'Transit Groups & Passengers' in the navigation menu.")

# 2. SHUTTLE TIMETABLE
elif nav_selection == "📅 Shuttle Timetable":
    st.markdown(
        """
        <style>
        .freeze-pane-container {
            max-height: 700px;
            overflow-y: auto;
            border: 1px solid #BFBFBF;
            border-radius: 5px;
            background-color: white;
            padding: 10px;
        }
        .freeze-pane-container table {
            width: 100%;
            border-collapse: collapse;
            background-color: white;
            color: black;
            font-family: Calibri, sans-serif;
            font-size: 14px;
        }
        .freeze-pane-container th {
            position: sticky;
            top: 0;
            background-color: #1F4E78;
            color: white;
            text-align: center;
            border: 1px solid #BFBFBF;
            padding: 10px;
            z-index: 5;
        }
        .freeze-pane-container td {
            border: 1px solid #BFBFBF;
            text-align: center;
            vertical-align: middle;
            padding: 8px;
            background-color: white;
        }
        </style>
        """,
        unsafe_allow_html=True
    )

    h_col1, h_col2, h_col3 = st.columns([1, 4, 1])
    with h_col1:
        if os.path.exists(LOGO1_PATH):
            st.image(LOGO1_PATH, width=130)
    with h_col2:
        st.markdown(
            """
            <div style="text-align: center;">
                <h3 style="color: #1F4E78; margin-bottom: 0px;">JGC SHUTTLE TIMETABLE</h3>
                <h4 style="color: #333333; margin-top: 2px; margin-bottom: 2px;">TUCC PROJECT - BATAM MODULE YARD [MD-1 & MD-4]</h4>
            </div>
            """,
            unsafe_allow_html=True
        )
    with h_col3:
        if os.path.exists(LOGO2_PATH):
            st.image(LOGO2_PATH, width=140)

    st.markdown("---")

    shuttle_raw_df = get_shuttle_raw_data()

    if not shuttle_raw_df.empty:
        shuttle_display_df = shuttle_raw_df.copy()
        shuttle_display_df['Days'] = "MONDAY<br>TUESDAY<br>WEDNESDAY<br>THURSDAY<br>FRIDAY<br>SATURDAY"
        shuttle_display_df['Unit'] = shuttle_display_df.apply(lambda r: f"{r['vehicle']} - {r['plate_number']}" if pd.notna(r['plate_number']) and r['plate_number'] != 'TBA' else "TOYOTA HI-ACE", axis=1)
        shuttle_display_df['Driver'] = shuttle_display_df['driver_name'].fillna("TBA")
        shuttle_display_df['Trip No.'] = [f"{i}st" if i==1 else f"{i}nd" if i==2 else f"{i}rd" if i==3 else f"{i}th" for i in range(1, len(shuttle_display_df)+1)]
        shuttle_display_df['Route (Yard-1)'] = shuttle_display_df['etd_1']
        shuttle_display_df['Route (Yard-3)'] = shuttle_display_df['etd_2']
        shuttle_display_df['Remarks'] = "DROP-OFF / PICK-UP"

        shuttle_rows_html = ""
        total_rows_count = len(shuttle_display_df)
        
        for idx, row in shuttle_display_df.iterrows():
            u_val = row['Unit']
            d_val = row['Driver']
            t_no = row['Trip No.']
            r_y1 = row['Route (Yard-1)']
            r_y3 = row['Route (Yard-3)']
            rem = row['Remarks']
            
            shuttle_rows_html += "<tr>"
            if idx == 0:
                shuttle_rows_html += f"<td rowspan='{total_rows_count}' style='font-weight: bold;'>MONDAY<br>TUESDAY<br>WEDNESDAY<br>THURSDAY<br>FRIDAY<br>SATURDAY</td>"
            
            shuttle_rows_html += f"<td>{u_val}</td>"
            shuttle_rows_html += f"<td>{d_val}</td>"
            shuttle_rows_html += f"<td>{t_no}</td>"
            shuttle_rows_html += f"<td>{r_y1}</td>"
            shuttle_rows_html += f"<td>{r_y3}</td>"
            shuttle_rows_html += f"<td>{rem}</td>"
            shuttle_rows_html += "</tr>"

        full_shuttle_html = f"""
        <div class="freeze-pane-container">
            <table>
                <thead>
                    <tr>
                        <th rowspan="2" style="top: 0; z-index: 6;">Days (s)</th>
                        <th rowspan="2" style="top: 0; z-index: 6;">UNIT</th>
                        <th rowspan="2" style="top: 0; z-index: 6;">DRIVER</th>
                        <th rowspan="2" style="top: 0; z-index: 6;">TRIP NO.</th>
                        <th colspan="2" style="top: 0; z-index: 6;">ROUTE</th>
                        <th rowspan="2" style="top: 0; z-index: 6;">REMARKS</th>
                    </tr>
                    <tr>
                        <th style="top: 41px; background-color: #245888; z-index: 5;">YARD - 1</th>
                        <th style="top: 41px; background-color: #245888; z-index: 5;">YARD - 3</th>
                    </tr>
                </thead>
                <tbody>
                    {shuttle_rows_html}
                </tbody>
            </table>
        </div>
        """
        st.markdown(full_shuttle_html, unsafe_allow_html=True)
        st.markdown("<br>", unsafe_allow_html=True)

        shuttle_excel_df = shuttle_raw_df.copy()
        shuttle_excel_df['Unit'] = shuttle_excel_df.apply(lambda r: f"{r['vehicle']} - {r['plate_number']}" if pd.notna(r['plate_number']) and r['plate_number'] != 'TBA' else "TOYOTA HI-ACE", axis=1)
        shuttle_excel_df['Driver'] = shuttle_excel_df['driver_name'].fillna("TBA")
        shuttle_excel_df['ETD Start'] = shuttle_excel_df['etd_1']
        shuttle_excel_df['ETD Return'] = shuttle_excel_df['etd_2']
        shuttle_excel_df['Remarks'] = "DROP-OFF"

        st.download_button(
            "📥 Download Shuttle Timetable Excel (.xlsx)",
            data=export_shuttle_timetable_excel(shuttle_excel_df),
            file_name=f"JGC_Shuttle_Timetable_{datetime.today().strftime('%Y%m%d')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        st.info("No scheduled transit dispatches found for Trip A, Trip B, or Trip C.")

# 3. OVERTIME & TRANSPORT
elif nav_selection == "⏰ Overtime & Transport":
    st.header("Request Overtime & Logistics Tracking")
    holiday_list = get_holidays_list()
    all_emp_names = get_employees_list()
    
    current_user_emp = st.session_state.get("emp_name", "") or st.session_state.get("username", "")
    if not all_emp_names:
        all_emp_names = [current_user_emp] if current_user_emp else ["Default Employee"]
    logged_in_emp = current_user_emp if current_user_emp in all_emp_names else all_emp_names[0]

    today_date = date.today()
    tomorrow_date = today_date + timedelta(days=1)
    tomorrow_is_sunday = tomorrow_date.weekday() == 6
    tomorrow_is_holiday = tomorrow_date.strftime("%Y-%m-%d") in holiday_list
    
    allowed_ot_dates = [today_date]
    if tomorrow_is_sunday or tomorrow_is_holiday:
        allowed_ot_dates.append(tomorrow_date)

    col1, col2 = st.columns(2)
    with col1:
        selected_staff_members = st.multiselect("Select Staff Member(s) for Overtime", options=all_emp_names, default=[logged_in_emp])
        ot_date = st.selectbox("Select Target Date", options=allowed_ot_dates, format_func=lambda d: d.strftime("%Y-%m-%d"))
        date_str = ot_date.strftime("%Y-%m-%d") if ot_date else ""
        
        is_sunday = ot_date.weekday() == 6 if ot_date else False
        is_holiday = date_str in holiday_list if date_str else False
        
        if is_sunday or is_holiday:
            default_start, default_end, default_origin, default_dest, default_dep_time_str = time(7, 0), time(15, 0), "Panbil", "Yard-1 Office", "07:00"
            st.caption("ℹ️ Baseline Rule: **Sunday/Holiday (07:00 - 15:00)**.")
        else:
            default_start, default_end, default_origin, default_dest, default_dep_time_str = time(17, 30), time(19, 0), "Yard-1 Office", "Panbil", "19:00"
            st.caption("ℹ️ Baseline Rule: **Weekday/Saturday (17:30 - 19:00)**.")
        
        start_time = st.time_input("OT Start Time (24-hr Military Time)", value=default_start)
        end_time = st.time_input("OT End Time (24-hr Military Time)", value=default_end)
        needs_transport = st.selectbox("Require Individual Transportation Logistics?", ["Yes", "No"], index=0)
        
    with col2:
        if needs_transport == "Yes":
            origin = st.text_input("Origin Address", value=default_origin)
            destination = st.text_input("Target Destination", value=default_dest)
            dep_time_str = st.text_input("Departure Timeline Estimate (24-hr)", value=default_dep_time_str)
            ret_time_str = st.text_input("Return Timeline Estimate (24-hr)", value="")
        else:
            origin, destination, dep_time_str, ret_time_str = ["", "", "", ""]

    if st.button("Submit New Overtime Request"):
        if not ot_date or not date_str:
            set_transaction_dialog("Data Transaction Unsuccessful", "Target Date cannot be left blank.", "error")
        elif not selected_staff_members:
            set_transaction_dialog("Data Transaction Unsuccessful", "Please select at least one staff member.", "error")
        else:
            try:
                conn = get_db_connection()
                cursor = conn.cursor()
                for staff in selected_staff_members:
                    cursor.execute('''
                        INSERT INTO overtime_requests (username, emp_name, ot_date, start_time, end_time, needs_transport, origin, destination, departure_time, return_time)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    ''', (st.session_state.username, staff, date_str, start_time.strftime("%H:%M"), end_time.strftime("%H:%M"), needs_transport, origin, destination, dep_time_str, ret_time_str))
                
                cursor.execute(
                    "SELECT id FROM daily_transit WHERE transit_date_start = %s AND transit_date_end = %s AND trip = 'Trip D'", 
                    (date_str, date_str)
                )
                existing_dispatch = cursor.fetchone()
                
                if not existing_dispatch:
                    cursor.execute('''
                        INSERT INTO daily_transit (transit_date_start, transit_date_end, group_name, requested_by, etd_1, etd_2, location_from, location_to, daily, trip)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    ''', (date_str, date_str, "TBA", st.session_state.username, dep_time_str if dep_time_str else "19:00", None, "Yard-1 Office", "Panbil", "No", "Trip D"))
                
                conn.commit()
                st.cache_data.clear()
                set_transaction_dialog("Data Transaction Successful", f"Overtime request logged for {len(selected_staff_members)} staff member(s) & transit dispatch schedule created.", "success")
            except Exception as e:
                set_transaction_dialog("Data Transaction Unsuccessful", f"Failed to save record: {str(e)}", "error")
        st.rerun()

    st.subheader("📋 Overtime Submission History Log")
    ot_df = get_overtime_requests_df(st.session_state.role, st.session_state.username)
    
    if not ot_df.empty:
        st.dataframe(ot_df, use_container_width=True)
        
        st.markdown("##### 📥 Export Current & Next Date Overtime Staff List")
        target_export_dates = [today_date.strftime("%Y-%m-%d")]
        if tomorrow_is_sunday or tomorrow_is_holiday:
            target_export_dates.append(tomorrow_date.strftime("%Y-%m-%d"))
            
        conn = get_db_connection()
        placeholders = ','.join(['%s'] * len(target_export_dates))
        summary_query = f"SELECT emp_name AS Employee_Name, ot_date AS Date, start_time AS Start_Time, end_time AS End_Time, needs_transport AS Needs_Transport, origin AS Origin, destination AS Destination FROM overtime_requests WHERE ot_date IN ({placeholders})"
        curr_next_ot_df = pd.read_sql_query(summary_query, conn, params=target_export_dates)
        
        if not curr_next_ot_df.empty:
            exp_col1, exp_col2 = st.columns(2)
            dates_label_str = " & ".join(target_export_dates)
            with exp_col1:
                st.download_button(
                    "📥 Export Current/Next OT Staff to Excel (.xlsx)",
                    data=export_overtime_summary_excel(curr_next_ot_df, title_text=f"Overtime Staff List ({dates_label_str})"),
                    file_name=f"Overtime_Staff_List_{datetime.today().strftime('%Y%m%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
            with exp_col2:
                if HAS_REPORTLAB:
                    st.download_button(
                        "📄 Export Current/Next OT Staff to PDF (.pdf)",
                        data=export_overtime_summary_pdf(curr_next_ot_df, title_text=f"Overtime Staff List ({dates_label_str})"),
                        file_name=f"Overtime_Staff_List_{datetime.today().strftime('%Y%m%d')}.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )

        if st.session_state.role in ["Admin", "Owner"]:
            with st.expander("✏️ Manage / Remove Overtime Submissions"):
                conn = get_db_connection()
                ot_raw = pd.read_sql_query("SELECT id, username, ot_date, emp_name, start_time, end_time, departure_time FROM overtime_requests ORDER BY id DESC", conn)
                
                if not ot_raw.empty:
                    today_str = date.today().strftime("%Y-%m-%d")
                    tomorrow_str = (date.today() + timedelta(days=1)).strftime("%Y-%m-%d")
                    
                    for _, o_row in ot_raw.iterrows():
                        o_id, o_uname, o_date, o_emp, o_start, o_end, o_dep = o_row['id'], o_row['username'], o_row['ot_date'], o_row['emp_name'], o_row['start_time'], o_row['end_time'], o_row['departure_time']
                        o_col1, o_col2, o_col3, o_col4 = st.columns([3, 2, 1, 1])
                        with o_col1:
                            st.text(f"ID #{o_id} | Date: {o_date} | Staff: {o_emp}")
                        with o_col2:
                            st.text(f"Time: {o_start} - {o_end}")
                        with o_col3:
                            if o_date in [today_str, tomorrow_str]:
                                if st.button("🚗 Transport", key=f"transport_ot_{o_id}", type="secondary"):
                                    try:
                                        conn = get_db_connection()
                                        cursor = conn.cursor()
                                        cursor.execute(
                                            "SELECT id FROM daily_transit WHERE transit_date_start = %s AND trip = 'Trip D'",
                                            (o_date,)
                                        )
                                        existing_dt = cursor.fetchone()
                                        
                                        if not existing_dt:
                                            cursor.execute('''
                                                INSERT INTO daily_transit (transit_date_start, transit_date_end, group_name, requested_by, etd_1, etd_2, location_from, location_to, daily, trip)
                                                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                                            ''', (o_date, o_date, "TBA", o_uname, o_dep if o_dep else "19:00", None, "Yard-1 Office", "Panbil", "No", "Trip D"))
                                            conn.commit()
                                            st.cache_data.clear()
                                            set_transaction_dialog("Transport Request Added", f"Daily transit dispatch (Trip D) created for date {o_date}.", "success")
                                        else:
                                            set_transaction_dialog("Already Exists", f"A Trip D transit dispatch for date {o_date} already exists.", "info")
                                    except Exception as e:
                                        set_transaction_dialog("Action Unsuccessful", f"Failed: {str(e)}", "error")
                                    st.rerun()
                            else:
                                st.text("Past date")
                        with o_col4:
                            if st.button("🗑️ Remove", key=f"del_ot_{o_id}", type="primary"):
                                try:
                                    conn = get_db_connection()
                                    cursor = conn.cursor()
                                    cursor.execute("DELETE FROM overtime_requests WHERE id = %s", (o_id,))
                                    conn.commit()
                                    st.cache_data.clear()
                                    set_transaction_dialog("Deletion Successful", f"Removed overtime record #{o_id}.", "success")
                                except Exception as e:
                                    set_transaction_dialog("Deletion Unsuccessful", f"Failed: {str(e)}", "error")
                                st.rerun()
                else:
                    st.info("No overtime submissions found.")

        st.download_button("📥 Export Overtime Log to Excel (.xlsx)", data=export_df_to_excel(ot_df, sheet_name="Overtime_Requests"), file_name=f"Overtime_Requests_{datetime.today().strftime('%Y%m%d')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# 4. TRANSIT GROUPS & PASSENGERS
elif nav_selection == "👥 Transit Groups & Passengers":
    st.header("👥 Transit Groups & Passengers Management")
    drivers_list = get_drivers_list()
    cars_list = get_cars_list()
    employees_list = get_employees_list()
    groups_list = get_transit_groups_list()

    col1, col2 = st.columns(2)
    with col1:
        st.subheader("1. Create Transit Group")
        selected_car_group = st.selectbox("Group Name (Select from Cars)", cars_list if cars_list else ["No cars available"], key="g_car_select")
        selected_driver = st.selectbox("Assign Driver (from Fleet Drivers)", drivers_list if drivers_list else ["No drivers available"])
        raw_etd1_grp = st.text_input("ETD 1 (From) [e.g. 0545 or 05:45]", value="05:45", key="grp_etd1_in")
        raw_etd2_grp = st.text_input("ETD 2 (To) [e.g. 1730 or 17:30]", value="17:30", key="grp_etd2_in")
        
        if st.button("Save Transit Group"):
            etd1_grp_formatted, etd2_grp_formatted = format_military_time(raw_etd1_grp), format_military_time(raw_etd2_grp)
            if not etd1_grp_formatted or not etd2_grp_formatted:
                set_transaction_dialog("Data Transaction Unsuccessful", "Invalid time format.", "error")
            else:
                try:
                    conn = get_db_connection()
                    cursor = conn.cursor()
                    cursor.execute("INSERT INTO transit_groups (group_name, driver_name, etd_1, etd_2) VALUES (%s, %s, %s, %s)", (selected_car_group, selected_driver, etd1_grp_formatted, etd2_grp_formatted))
                    conn.commit()
                    st.cache_data.clear()
                    set_transaction_dialog("Data Transaction Successful", f"Group '{selected_car_group}' created.", "success")
                except Exception as e:
                    set_transaction_dialog("Data Transaction Unsuccessful", f"Failed: {str(e)}", "error")
            st.rerun()

    with col2:
        st.subheader("2. Assign Passengers to Group")
        if groups_list and employees_list:
            selected_group_for_p = st.selectbox("Select Group Name", groups_list, key="p_group_select")
            selected_passengers = st.multiselect("Select Passenger(s)", employees_list)
            if st.button("Assign Passengers to Group"):
                if selected_passengers:
                    try:
                        conn = get_db_connection()
                        cursor = conn.cursor()
                        for p in selected_passengers:
                            cursor.execute("INSERT INTO transit_passengers (group_name, passengers) VALUES (%s, %s)", (selected_group_for_p, p))
                        conn.commit()
                        st.cache_data.clear()
                        set_transaction_dialog("Data Transaction Successful", "Passengers assigned.", "success")
                    except Exception as e:
                        set_transaction_dialog("Data Transaction Unsuccessful", f"Failed: {str(e)}", "error")
                st.rerun()

    st.markdown("---")
    title_col, eff_date_col = st.columns([2, 1])
    with title_col:
        st.subheader("📋 Configured Groups & Assigned Passengers")
    with eff_date_col:
        target_effective_date = st.date_input("Target Effective Date", value=date.today(), key="eff_date_picker")
        eff_date_str = target_effective_date.strftime("%Y-%m-%d") if target_effective_date else ""

    unrolled_df = get_home_unrolled_data()
    
    if not unrolled_df.empty:
        st.dataframe(unrolled_df, use_container_width=True)
        
        with st.expander("✏️ Manage / Remove Passenger Assignments"):
            conn = get_db_connection()
            passengers_raw = pd.read_sql_query("SELECT id, group_name, passengers FROM transit_passengers", conn)
            
            if not passengers_raw.empty:
                for _, p_row in passengers_raw.iterrows():
                    p_id, p_grp, p_name = p_row['id'], p_row['group_name'], p_row['passengers']
                    p_col1, p_col2 = st.columns([4, 1])
                    with p_col1:
                        st.text(f"Group: {p_grp} | Passenger: {p_name}")
                    with p_col2:
                        if st.button("🗑️ Remove", key=f"del_passenger_{p_id}", type="primary"):
                            try:
                                conn = get_db_connection()
                                cursor = conn.cursor()
                                cursor.execute("DELETE FROM transit_passengers WHERE id = %s", (p_id,))
                                conn.commit()
                                st.cache_data.clear()
                                set_transaction_dialog("Deletion Successful", f"Passenger assignment #{p_id} removed.", "success")
                            except Exception as e:
                                set_transaction_dialog("Deletion Unsuccessful", f"Failed: {str(e)}", "error")
                            st.rerun()

        st.markdown("##### 📥 Export Passenger List")
        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            st.download_button(
                "📥 Export Passenger List to Excel (.xlsx)",
                data=export_custom_batam_excel(unrolled_df, effective_date_str=eff_date_str),
                file_name=f"Daily_Transportation_Arrangement_{datetime.today().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        with dl_col2:
            if HAS_REPORTLAB:
                st.download_button(
                    "📄 Export Passenger List to PDF (.pdf)",
                    data=export_custom_batam_pdf(unrolled_df, effective_date_str=eff_date_str),
                    file_name=f"Daily_Transportation_Arrangement_{datetime.today().strftime('%Y%m%d')}.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )

# 5. DAILY TRANSIT DISPATCH SETUP
elif nav_selection == "📅 Daily Transit Dispatch Setup":
    st.header("📅 Daily Transit Dispatch Schedule Setup")
    
    with st.form("transit_dispatch_form"):
        col1, col2 = st.columns(2)
        with col1:
            start_date = st.date_input("Start Date", value=date.today())
            end_date = st.date_input("End Date", value=date.today())
            
            car_names = get_cars_list()
            trip_options = get_trips_list()
            
            group_name = st.selectbox("Assign Car / Group", options=car_names if car_names else ["TBA"])
            trip_type = st.selectbox("Trip Schedule Category", options=trip_options if trip_options else ["Trip A"])
            daily_repeat = st.selectbox("Repeat Daily?", options=["No", "Yes"])
            
        with col2:
            raw_etd1 = st.text_input("ETD 1 (From) [e.g. 0545 or 05:45]", value="05:45")
            raw_etd2 = st.text_input("ETD 2 (To) [e.g. 1730 or 17:30]", value="17:30")
            loc_from = st.text_input("Origin Location", value="Yard-1 Office")
            loc_to = st.text_input("Destination Location", value="Panbil")
            
        submitted_transit = st.form_submit_button("Create Transit Dispatch Entry")
        if submitted_transit:
            etd1_fmt, etd2_fmt = format_military_time(raw_etd1), format_military_time(raw_etd2)
            if not etd1_fmt or not etd2_fmt:
                set_transaction_dialog("Data Transaction Unsuccessful", "Invalid ETD time format.", "error")
            else:
                try:
                    conn = get_db_connection()
                    cursor = conn.cursor()
                    cursor.execute('''
                        INSERT INTO daily_transit (transit_date_start, transit_date_end, group_name, requested_by, etd_1, etd_2, location_from, location_to, daily, trip)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    ''', (start_date.strftime("%Y-%m-%d"), end_date.strftime("%Y-%m-%d"), group_name, st.session_state.username, etd1_fmt, etd2_fmt, loc_from, loc_to, daily_repeat, trip_type))
                    conn.commit()
                    st.cache_data.clear()
                    set_transaction_dialog("Data Transaction Successful", "Transit dispatch schedule successfully logged.", "success")
                except Exception as e:
                    set_transaction_dialog("Data Transaction Unsuccessful", f"Failed: {str(e)}", "error")
            st.rerun()

    st.subheader("📋 Active Daily Transit Dispatches")
    transit_df = get_daily_transit_df()
    
    if not transit_df.empty:
        st.dataframe(transit_df, use_container_width=True)
        
        if st.session_state.role in ["Admin", "Owner"]:
            with st.expander("✏️ Manage / Remove Transit Dispatches"):
                conn = get_db_connection()
                t_raw = pd.read_sql_query("SELECT id, transit_date_start, group_name, trip FROM daily_transit ORDER BY id DESC", conn)
                for _, tr_row in t_raw.iterrows():
                    tr_id, tr_date, tr_grp, tr_trip = tr_row['id'], tr_row['transit_date_start'], tr_row['group_name'], tr_row['trip']
                    tc1, tc2 = st.columns([4, 1])
                    with tc1:
                        st.text(f"ID #{tr_id} | Date: {tr_date} | Car: {tr_grp} | Trip: {tr_trip}")
                    with tc2:
                        if st.button("🗑️ Remove", key=f"del_transit_{tr_id}", type="primary"):
                            try:
                                conn = get_db_connection()
                                cursor = conn.cursor()
                                cursor.execute("DELETE FROM daily_transit WHERE id = %s", (tr_id,))
                                conn.commit()
                                st.cache_data.clear()
                                set_transaction_dialog("Deletion Successful", f"Removed transit dispatch #{tr_id}.", "success")
                            except Exception as e:
                                set_transaction_dialog("Deletion Unsuccessful", f"Failed: {str(e)}", "error")
                            st.rerun()

        st.download_button("📥 Export Daily Transit to Excel (.xlsx)", data=export_df_to_excel(transit_df, sheet_name="Daily_Transit"), file_name=f"Daily_Transit_{datetime.today().strftime('%Y%m%d')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# 6. MEETING ROOMS
elif nav_selection == "🏢 Meeting Rooms":
    st.header("🏢 Meeting Room Booking Portal")
    rooms_df = get_meeting_rooms_df()
    
    st.subheader("Available Meeting Rooms")
    st.dataframe(rooms_df, use_container_width=True)
    
    st.markdown("---")
    st.subheader("Book a Meeting Room")
    with st.form("room_booking_form"):
        r_nums = get_room_numbers_list()
        
        selected_room = st.selectbox("Select Room Number", options=r_nums if r_nums else ["101"])
        booking_date = st.date_input("Booking Date", value=date.today())
        start_t = st.time_input("Start Time", value=time(9, 0))
        end_t = st.time_input("End Time", value=time(10, 0))
        is_recurring = st.selectbox("Is Recurring Weekly?", ["No", "Yes"])
        rec_end_date = st.date_input("Recurrence End Date (if recurring)", value=date.today())
        
        submitted_booking = st.form_submit_button("Confirm Room Booking")
        if submitted_booking:
            try:
                conn = get_db_connection()
                cursor = conn.cursor()
                cursor.execute('''
                    INSERT INTO room_bookings (room_number, booked_by, booking_date, start_time, end_time, is_recurring, recurrence_end_date)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                ''', (selected_room, st.session_state.username, booking_date.strftime("%Y-%m-%d"), start_t.strftime("%H:%M"), end_t.strftime("%H:%M"), is_recurring, rec_end_date.strftime("%Y-%m-%d") if is_recurring == "Yes" else ""))
                conn.commit()
                st.cache_data.clear()
                set_transaction_dialog("Booking Successful", f"Meeting room {selected_room} booked successfully.", "success")
            except Exception as e:
                set_transaction_dialog("Booking Unsuccessful", f"Failed: {str(e)}", "error")
            st.rerun()

    st.subheader("📋 Active Room Bookings")
    bookings_df = get_room_bookings_df()
    
    if not bookings_df.empty:
        st.dataframe(bookings_df, use_container_width=True)
        if st.session_state.role in ["Admin", "Owner"]:
            with st.expander("✏️ Manage / Remove Room Bookings"):
                conn = get_db_connection()
                b_raw = pd.read_sql_query("SELECT id, room_number, booking_date FROM room_bookings", conn)
                for _, b_row in b_raw.iterrows():
                    b_id, b_room, b_date = b_row['id'], b_row['room_number'], b_row['booking_date']
                    bc1, bc2 = st.columns([4, 1])
                    with bc1:
                        st.text(f"ID #{b_id} | Room: {b_room} | Date: {b_date}")
                    with bc2:
                        if st.button("🗑️ Remove", key=f"del_booking_{b_id}", type="primary"):
                            try:
                                conn = get_db_connection()
                                cursor = conn.cursor()
                                cursor.execute("DELETE FROM room_bookings WHERE id = %s", (b_id,))
                                conn.commit()
                                st.cache_data.clear()
                                set_transaction_dialog("Deletion Successful", f"Removed booking #{b_id}.", "success")
                            except Exception as e:
                                set_transaction_dialog("Deletion Unsuccessful", f"Failed: {str(e)}", "error")
                            st.rerun()

# 7. SITE NEWS
elif nav_selection == "📢 Site News":
    st.header("📢 Site News & Announcements")
    
    if st.session_state.role in ["Admin", "Owner"]:
        with st.form("upload_news_form"):
            news_title = st.text_input("Announcement Title")
            uploaded_file = st.file_uploader("Upload Document / Notice (PDF, Excel, Images)", type=["pdf", "xlsx", "xls", "png", "jpg", "jpeg"])
            submitted_news = st.form_submit_button("Publish Announcement")
            
            if submitted_news:
                if not news_title or not uploaded_file:
                    set_transaction_dialog("Publication Unsuccessful", "Title and file upload cannot be blank.", "error")
                else:
                    try:
                        file_path = os.path.join(NEWS_DIR, uploaded_file.name)
                        with open(file_path, "wb") as f:
                            f.write(uploaded_file.getbuffer())
                        
                        conn = get_db_connection()
                        cursor = conn.cursor()
                        cursor.execute('''
                            INSERT INTO site_news (title, filename, file_path, file_type, uploaded_by, upload_date)
                            VALUES (%s, %s, %s, %s, %s, %s)
                        ''', (news_title, uploaded_file.name, file_path, uploaded_file.type, st.session_state.username, datetime.today().strftime("%Y-%m-%d")))
                        conn.commit()
                        st.cache_data.clear()
                        set_transaction_dialog("Publication Successful", "Site news published successfully.", "success")
                    except Exception as e:
                        set_transaction_dialog("Publication Unsuccessful", f"Failed: {str(e)}", "error")
                st.rerun()

    news_df = get_site_news_df()
    
    if not news_df.empty:
        for _, n_row in news_df.iterrows():
            n_id, n_title, n_fname, n_fpath, n_author, n_date = n_row['id'], n_row['title'], n_row['filename'], n_row['file_path'], n_row['uploaded_by'], n_row['upload_date']
            with st.container():
                st.subheader(n_title)
                st.caption(f"Published by: {n_author} on {n_date}")
                if os.path.exists(n_fpath):
                    with open(n_fpath, "rb") as f:
                        file_bytes = f.read()
                    st.download_button(f"📥 Download {n_fname}", data=file_bytes, file_name=n_fname, key=f"dl_news_{n_id}")
                if st.session_state.role in ["Admin", "Owner"]:
                    if st.button("🗑️ Delete Announcement", key=f"del_news_{n_id}", type="secondary"):
                        try:
                            if os.path.exists(n_fpath):
                                os.remove(n_fpath)
                            conn = get_db_connection()
                            cursor = conn.cursor()
                            cursor.execute("DELETE FROM site_news WHERE id = %s", (n_id,))
                            conn.commit()
                            st.cache_data.clear()
                            set_transaction_dialog("Deletion Successful", "Announcement removed.", "success")
                        except Exception as e:
                            set_transaction_dialog("Deletion Unsuccessful", f"Failed: {str(e)}", "error")
                        st.rerun()
                st.markdown("---")
    else:
        st.info("No site news announcements published yet.")

# 8. SYSTEM ADMINISTRATION
elif nav_selection == "🛠️ System Administration":
    st.header("🛠️ System Administration & Master Data Management")
    
    if st.session_state.role not in ["Admin", "Owner"]:
        st.error("Access Denied. Administrator privileges required to access System Administration.")
    else:
        tab_backup, tab_import, tab_master = st.tabs(["💾 Backup & Recovery", "📥 Excel Template & Smart Import", "⚙ Master Tables Editor"])
        
        with tab_backup:
            st.subheader("Database Backup & Export")
            st.markdown("Create a manual timestamped Excel backup export of all PostgreSQL tables.")
            if st.button("Perform Manual Backup Now", type="primary"):
                success, db_path, excel_path = perform_manual_backup()
                if success:
                    set_transaction_dialog("Backup Successful", f"Excel export saved to: {excel_path}", "success")
                else:
                    set_transaction_dialog("Backup Failed", f"Error: {db_path}", "error")
                st.rerun()
                
            st.markdown("##### Existing Backup Files:")
            if os.path.exists(BACKUP_DIR):
                backup_files = sorted(os.listdir(BACKUP_DIR), reverse=True)
                if backup_files:
                    for b_file in backup_files[:10]:
                        f_path = os.path.join(BACKUP_DIR, b_file)
                        with open(f_path, "rb") as bf:
                            b_data = bf.read()
                        st.download_button(f"📥 Download {b_file}", data=b_data, file_name=b_file, key=f"dl_backup_{b_file}")
                else:
                    st.info("No backup files found.")

        with tab_import:
            st.subheader("📥 Excel Template Download & Smart Data Import")
            st.markdown(
                """
                Download the master Excel template containing all database tables as sheets. 
                Fill or update your data in the respective sheets, and upload it back. 
                The importer is smart: it matches columns and **automatically appends new rows or updates existing records** based on each table's unique key!
                """
            )
            
            # 1. GENERATE & DOWNLOAD TEMPLATE
            conn = get_db_connection()
            table_names = [
                "users", "holidays", "trips", "overtime_requests", "meeting_rooms", 
                "room_bookings", "cars", "fleet_drivers", "transit_groups", 
                "transit_passengers", "daily_transit", "site_news"
            ]
            
            template_buffer = io.BytesIO()
            with pd.ExcelWriter(template_buffer, engine='openpyxl') as writer:
                for tname in table_names:
                    try:
                        df_tbl = pd.read_sql_query(f"SELECT * FROM {tname} LIMIT 0", conn)
                    except Exception:
                        df_tbl = pd.DataFrame()
                    df_tbl.to_excel(writer, index=False, sheet_name=tname)
            
            st.download_button(
                "📥 Download Master Excel Import Template (.xlsx)",
                data=template_buffer.getvalue(),
                file_name=f"Master_Operations_Import_Template_{datetime.today().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
            
            st.markdown("---")
            st.subheader("Upload & Process Master Excel Import")
            uploaded_import_file = st.file_uploader("Upload Completed Excel Master Template", type=["xlsx", "xls"])
            
            if uploaded_import_file is not None:
                if st.button("🚀 Process Smart Import & Update Database", type="primary"):
                    try:
                        excel_file_obj = pd.ExcelFile(uploaded_import_file)
                        conn = get_db_connection()
                        cursor = conn.cursor()
                        
                        import_summary_log = []
                        
                        pk_mapping = {
                            "users": "username",
                            "holidays": "holiday_date",
                            "trips": "trip",
                            "overtime_requests": "id",
                            "meeting_rooms": "room_number",
                            "room_bookings": "id",
                            "cars": "car_name",
                            "fleet_drivers": "driver_name",
                            "transit_groups": "group_name",
                            "transit_passengers": "id",
                            "daily_transit": "id",
                            "site_news": "id"
                        }
                        
                        for sheet_name in excel_file_obj.sheet_names:
                            if sheet_name in table_names:
                                df_sheet = pd.read_excel(excel_file_obj, sheet_name=sheet_name)
                                if df_sheet.empty:
                                    continue
                                    
                                pk_col = pk_mapping.get(sheet_name)
                                updated_count = 0
                                inserted_count = 0
                                
                                for _, row in df_sheet.iterrows():
                                    clean_row = row.dropna()
                                    if clean_row.empty:
                                        continue
                                        
                                    cols = list(clean_row.index)
                                    vals = list(clean_row.values)
                                    
                                    insert_cols = cols
                                    insert_vals = vals
                                    if pk_col in cols and sheet_name in ["overtime_requests", "room_bookings", "transit_passengers", "daily_transit", "site_news"]:
                                        if pd.isna(clean_row[pk_col]) or str(clean_row[pk_col]).strip() == "":
                                            insert_cols = [c for c in cols if c != pk_col]
                                            insert_vals = [clean_row[c] for c in insert_cols]
                                            
                                    placeholders = ", ".join(["%s"] * len(insert_cols))
                                    col_names_str = ", ".join(insert_cols)
                                    
                                    if pk_col and pk_col in insert_cols:
                                        update_set_clause = ", ".join([f"{c} = EXCLUDED.{c}" for c in insert_cols if c != pk_col])
                                        if update_set_clause:
                                            sql = f"INSERT INTO {sheet_name} ({col_names_str}) VALUES ({placeholders}) ON CONFLICT ({pk_col}) DO UPDATE SET {update_set_clause}"
                                        else:
                                            sql = f"INSERT INTO {sheet_name} ({col_names_str}) VALUES ({placeholders}) ON CONFLICT ({pk_col}) DO NOTHING"
                                    else:
                                        sql = f"INSERT INTO {sheet_name} ({col_names_str}) VALUES ({placeholders})"
                                        
                                    cursor.execute(sql, insert_vals)
                                    inserted_count += 1
                                    
                                import_summary_log.append(f"Table **{sheet_name}**: Processed {inserted_count} records")
                                
                        conn.commit()
                        st.cache_data.clear()
                        
                        summary_msg = "\n".join(import_summary_log)
                        set_transaction_dialog("Smart Import Successful", f"Database updated successfully!\n\n{summary_msg}", "success")
                    except Exception as e:
                        set_transaction_dialog("Import Failed", f"Error during import processing: {str(e)}", "error")
                    st.rerun()

        with tab_master:
            st.subheader("Master Tables Data View & Quick Management")
            selected_master_table = st.selectbox(
                "Select Master Table to Inspect",
                ["users", "cars", "fleet_drivers", "transit_groups", "transit_passengers", "meeting_rooms", "holidays", "trips"]
            )
            
            conn = get_db_connection()
            master_df = pd.read_sql_query(f"SELECT * FROM {selected_master_table}", conn)
            
            st.dataframe(master_df, use_container_width=True)
            st.download_button(
                f"📥 Export {selected_master_table} to Excel (.xlsx)",
                data=export_df_to_excel(master_df, sheet_name=selected_master_table),
                file_name=f"{selected_master_table}_{datetime.today().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            
            st.markdown("---")
            st.subheader(f"Manage Records (Grid Editor): `{selected_master_table}`")
            st.markdown("Edit cells directly, add new rows at the bottom, or select rows to delete. Click **Save Changes to Database** when done.")

            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute("""
                SELECT kcu.column_name 
                FROM information_schema.table_constraints tc 
                JOIN information_schema.key_column_usage kcu ON tc.constraint_name = kcu.constraint_name 
                WHERE tc.table_name = %s AND tc.constraint_type = 'PRIMARY KEY'
            """, (selected_master_table,))
            pk_rows = cursor.fetchall()
            pk_cols = [row[0] for row in pk_rows]
            if not pk_cols:
                pk_cols = [master_df.columns[0] if not master_df.empty else '']
            pk_col = pk_cols[0] if pk_cols else None

            edited_df = st.data_editor(
                master_df,
                num_rows="dynamic",
                use_container_width=True,
                key=f"grid_editor_{selected_master_table}"
            )

            if st.button("💾 Save Changes to Database", type="primary", key=f"save_grid_{selected_master_table}"):
                try:
                    conn = get_db_connection()
                    cursor = conn.cursor()
                    
                    if pk_col and pk_col in edited_df.columns and pk_col in master_df.columns:
                        original_keys = set(master_df[pk_col].dropna().astype(str))
                        current_keys = set(edited_df[pk_col].dropna().astype(str))
                        
                        deleted_keys = original_keys - current_keys
                        for d_key in deleted_keys:
                            cursor.execute(f"DELETE FROM {selected_master_table} WHERE CAST({pk_col} AS TEXT) = %s", (d_key,))
                            
                        for _, row in edited_df.iterrows():
                            clean_row = row.dropna()
                            if clean_row.empty:
                                continue
                            cols = list(clean_row.index)
                            p_val = str(row[pk_col]) if pk_col in row and pd.notna(row[pk_col]) else ""
                            
                            if p_val and p_val in original_keys:
                                set_cols = [c for c in cols if c != pk_col]
                                set_clause = ", ".join([f"{c} = %s" for c in set_cols])
                                update_vals = [row[c] for c in set_cols] + [row[pk_col]]
                                if set_clause:
                                    cursor.execute(f"UPDATE {selected_master_table} SET {set_clause} WHERE {pk_col} = %s", update_vals)
                            else:
                                insert_cols = [c for c in cols if not (pk_col and c == pk_col and (pd.isna(row[c]) or str(row[c]).strip() == ""))]
                                insert_vals = [row[c] for c in insert_cols]
                                placeholders = ", ".join(["%s"] * len(insert_cols))
                                col_names_str = ", ".join(insert_cols)
                                if insert_cols:
                                    if pk_col and pk_col in insert_cols:
                                        update_set_clause = ", ".join([f"{c} = EXCLUDED.{c}" for c in insert_cols if c != pk_col])
                                        sql = f"INSERT INTO {selected_master_table} ({col_names_str}) VALUES ({placeholders}) ON CONFLICT ({pk_col}) DO UPDATE SET {update_set_clause}"
                                    else:
                                        sql = f"INSERT INTO {selected_master_table} ({col_names_str}) VALUES ({placeholders})"
                                    cursor.execute(sql, insert_vals)
                    else:
                        cursor.execute(f"DELETE FROM {selected_master_table}")
                        for _, row in edited_df.iterrows():
                            clean_row = row.dropna()
                            if clean_row.empty:
                                continue
                            cols = list(clean_row.index)
                            vals = list(clean_row.values)
                            placeholders = ", ".join(["%s"] * len(cols))
                            col_names_str = ", ".join(cols)
                            cursor.execute(f"INSERT INTO {selected_master_table} ({col_names_str}) VALUES ({placeholders})", vals)
                            
                    conn.commit()
                    st.cache_data.clear()
                    set_transaction_dialog("Save Successful", f"Master table `{selected_master_table}` updated successfully.", "success")
                except Exception as e:
                    set_transaction_dialog("Save Failed", f"Error: {str(e)}", "error")
                st.rerun()
